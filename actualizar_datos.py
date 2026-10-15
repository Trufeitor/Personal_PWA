#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de Actualización y Cifrado para Personal_PWA
Lee las plantillas Excel de Personal_XT, realiza la unión de datos (Centros y Categorías)
y cifra el resultado en data.js usando AES-256-GCM para máxima seguridad en GitHub Pages.
"""

import os
import sys
import glob
import json
import base64
import argparse
import datetime
import openpyxl
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

def normalize_text(text):
    if not text:
        return ""
    import unicodedata
    normalized = unicodedata.normalize("NFD", str(text))
    no_accents = "".join(c for c in normalized if unicodedata.category(c) != "Mn")
    import re
    cleaned = re.sub(r"[^a-zA-Z0-9]", " ", no_accents).lower()
    return re.sub(r"\s+", " ", cleaned).strip()

def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()

def clean_date(val):
    if val is None:
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return None
    parts = s.split("/")
    if len(parts) == 3:
        d = parts[0].zfill(2)
        m = parts[1].zfill(2)
        y = parts[2]
        return f"{y}-{m}-{d}"
    return s

def find_file(pattern, search_dirs):
    for d in search_dirs:
        matches = glob.glob(os.path.join(d, pattern))
        if matches:
            return matches[0]
    return None

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    local_config_file = os.path.join(script_dir, "config.local.json")
    saved_pin = None
    if os.path.exists(local_config_file):
        try:
            with open(local_config_file, "r", encoding="utf-8") as f:
                cfg = json.load(f)
                saved_pin = cfg.get("pin")
        except Exception:
            pass

    parser = argparse.ArgumentParser(description="Actualizar y cifrar datos para Personal_PWA")
    parser.add_argument("--pin", default=saved_pin, help="PIN de desbloqueo (si no se especifica, se lee de config.local.json)")
    parser.add_argument("--src", default=r"C:\Users\HTPC\Documents\GitHub\Personal_XT", help="Carpeta origen con los Excel")
    args = parser.parse_args()

    pin = args.pin
    if not pin:
        import getpass
        pin = getpass.getpass("Introduce el PIN de cifrado: ").strip()
    if not pin:
        print("Error: Se requiere un PIN de cifrado.")
        sys.exit(1)

    search_dirs = [args.src, ".", "..", os.path.join("..", "Personal_XT")]

    file_plantilla = find_file("Plantilla*.xlsx", search_dirs)
    file_categorias = find_file("Categorias*.xlsx", search_dirs)
    file_centros = find_file("*Centros*.xlsx", search_dirs)

    if not file_plantilla or not file_categorias or not file_centros:
        print(f"Error: No se encontraron todos los archivos Excel necesarios en {search_dirs}")
        print(f"Plantilla: {file_plantilla}")
        print(f"Categorías: {file_categorias}")
        print(f"Centros: {file_centros}")
        sys.exit(1)

    print(f"Cargando Categorías desde: {file_categorias}")
    wb_cat = openpyxl.load_workbook(file_categorias, data_only=True)
    sheet_cat = wb_cat.active
    categorias = {}
    for row in sheet_cat.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            categorias[clean_str(row[0])] = clean_str(row[1])

    print(f"Cargando Centros desde: {file_centros}")
    wb_cen = openpyxl.load_workbook(file_centros, data_only=True)
    sheet_cen = wb_cen.active
    centros = {}
    for row in sheet_cen.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            codigo = clean_str(row[0]).zfill(4)
            centros[codigo] = clean_str(row[1])

    print(f"Cargando Plantilla de Personal desde: {file_plantilla}")
    wb_pla = openpyxl.load_workbook(file_plantilla, data_only=True)
    sheet_pla = wb_pla.active

    empleados = []
    for row in sheet_pla.iter_rows(min_row=2, values_only=True):
        if not row or row[1] is None:
            continue
        try:
            rubrica = int(row[1])
        except (ValueError, TypeError):
            continue

        codigo_centro = clean_str(row[0]).zfill(4)
        nombre = clean_str(row[2])
        codigo_cat = clean_str(row[3])
        nivel = clean_str(row[4])
        ant = clean_date(row[5])
        ant_cat = clean_date(row[6])
        nac = clean_date(row[7])
        raw_tipo = clean_str(row[8]).upper()
        tipo = "T" if raw_tipo.startswith("T") else "I"

        nom_cat = categorias.get(codigo_cat, "")
        nom_cen = centros.get(codigo_centro, "")

        if nombre:
            empleados.append({
                "r": rubrica,
                "n": nombre,
                "nv": nivel,
                "cc": codigo_centro,
                "nc": nom_cen,
                "ct": codigo_cat,
                "nct": nom_cat,
                "tc": tipo,
                "a": ant,
                "ac": ant_cat,
                "fn": nac,
                # Campos pre-normalizados para búsqueda ultrarrápida
                "nb": normalize_text(nombre),
                "nvb": normalize_text(nivel),
                "ncb": normalize_text(nom_cen),
                "nctb": normalize_text(nom_cat)
            })

    # Ordenar alfabéticamente por nombre
    empleados.sort(key=lambda x: x["n"])

    print(f"Total empleados procesados con éxito: {len(empleados)}")

    # Serializar a JSON
    payload = {
        "version": "1.0",
        "fecha": datetime.datetime.now().strftime("%d/%m/%Y"),
        "total": len(empleados),
        "data": empleados
    }
    plaintext = json.dumps(payload, ensure_ascii=False).encode("utf-8")

    # Cifrado AES-256-GCM con PBKDF2
    print(f"Cifrando datos ({len(empleados)} registros, clave de {len(pin)} dígitos)...")
    salt = os.urandom(16)
    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(),
        length=32,
        salt=salt,
        iterations=100000,
    )
    key = kdf.derive(pin.encode("utf-8"))
    aesgcm = AESGCM(key)
    iv = os.urandom(12)
    ciphertext = aesgcm.encrypt(iv, plaintext, None)

    output_data = {
        "salt": base64.b64encode(salt).decode("ascii"),
        "iv": base64.b64encode(iv).decode("ascii"),
        "ciphertext": base64.b64encode(ciphertext).decode("ascii"),
        "total": len(empleados),
        "pin_length": len(pin),
        "actualizado": datetime.datetime.now().strftime("%d/%m/%Y")
    }

    out_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.js")
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("// Archivo generado automáticamente. Datos cifrados con AES-256-GCM.\n")
        f.write(f"window.ENCRYPTED_DATA = {json.dumps(output_data, indent=2)};\n")

    print(f"Archivo guardado exitosamente: {out_file}")
    print("¡Listo! El paquete de datos está cifrado e integrado.")

if __name__ == "__main__":
    main()
